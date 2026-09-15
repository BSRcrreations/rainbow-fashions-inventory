"""Read-only, store-scoped business workbook and per-size replenishment list."""
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO
from uuid import UUID

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy import func
from sqlalchemy.orm import joinedload
from app.core.exceptions import bad_request
from app.models import Product, ProductVariant, Sale, SaleItem, SaleReturn, SaleReturnItem, Purchase, StockHistory, InventoryCostLot
from app.models.purchase_item import PurchaseItem
from app.models.purchase_return import PurchaseReturn
from app.models.customer import Customer, CustomerPayment
from app.models.supplier import Supplier, SupplierPayment
from app.models.expense import Expense
from app.models.operations import DayClosing
from app.models.enums import SaleStatus, PurchaseStatus
from app.services.business_service import SupplierService, CustomerService
from app.services.operations_service import business_day_bounds, get_store_settings
from app.services.report_service import ReportService


class ReportExportService:
    def __init__(self, db): self.db = db

    def low_stock_query(self, user, category_id=None, brand_id=None, supplier_id=None):
        if not user.store_id: raise bad_request("A store is required.")
        query = self.db.query(ProductVariant).join(Product, Product.id == ProductVariant.product_id).options(joinedload(ProductVariant.product).joinedload(Product.category), joinedload(ProductVariant.product).joinedload(Product.brand)).filter(ProductVariant.store_id == user.store_id, Product.store_id == user.store_id, ProductVariant.is_active.is_(True), Product.is_active.is_(True), ProductVariant.current_stock <= func.coalesce(ProductVariant.minimum_stock, Product.minimum_stock))
        if category_id: query = query.filter(Product.category_id == category_id)
        if brand_id: query = query.filter(Product.brand_id == brand_id)
        if supplier_id:
            received = self.db.query(PurchaseItem.product_variant_id).join(Purchase, Purchase.id == PurchaseItem.purchase_id).filter(Purchase.store_id == user.store_id, Purchase.supplier_id == supplier_id, Purchase.status == PurchaseStatus.CONFIRMED)
            query = query.filter(ProductVariant.id.in_(received))
        return query.order_by(ProductVariant.current_stock, Product.name, ProductVariant.id)

    def low_stock(self, user, page=1, category_id=None, brand_id=None, supplier_id=None):
        rows = self.low_stock_query(user, category_id, brand_id, supplier_id).offset((page-1)*50).limit(51).all()
        return {"page":page, "has_more":len(rows)>50, "items":[self.stock_row(row) for row in rows[:50]]}

    @staticmethod
    def stock_row(variant):
        product=variant.product
        return {"product":product.name,"category":product.category.name,"brand":product.brand.name,"size":variant.size,"color":variant.color,"sku":variant.internal_sku,"barcode":variant.barcode,"current_stock":variant.current_stock,"minimum_stock":variant.minimum_stock if variant.minimum_stock is not None else product.minimum_stock}

    @staticmethod
    def limited(query):
        rows=query.limit(50001).all()
        if len(rows)>50000: raise bad_request("This report is too large. Choose a shorter date range.")
        return rows

    def workbook(self, user, start_date: date, end_date: date):
        if start_date>end_date: raise bad_request("End date cannot be earlier than Start date.")
        if (end_date-start_date).days>366: raise bad_request("Choose a report period of at most one year.")
        store_id=user.store_id
        if not store_id: raise bad_request("A store is required.")
        zone=get_store_settings(self.db,store_id).timezone
        start,_=business_day_bounds(start_date,zone); _,end=business_day_bounds(end_date,zone)
        book=Workbook();book.remove(book.active)
        def sheet(title,headers,rows):
            ws=book.create_sheet(title);ws.append(headers);ws.freeze_panes="A2"
            for cell in ws[1]:cell.font=Font(bold=True,color="FFFFFF");cell.fill=PatternFill("solid",fgColor="0F766E")
            for row in rows:
                ws.append([float(value) if isinstance(value,Decimal) else ("'"+value if isinstance(value,str) and value.startswith(("=","+","-","@")) else value) for value in row])
            ws.auto_filter.ref=ws.dimensions
            for column in ws.columns:ws.column_dimensions[column[0].column_letter].width=min(45,max(14,max(len(str(cell.value or "")) for cell in column)+2))
        summary=ReportService(self.db).summary(user,start_date,end_date)
        sheet("Summary",["Period / measure","Value"],[["From",start_date.isoformat()],["To",end_date.isoformat()],["Timezone",zone],["Net sales",summary.profit_and_loss.sales_total],["Gross profit",summary.profit_and_loss.gross_profit],["Expenses",summary.profit_and_loss.expense_total],["Net profit",summary.profit_and_loss.net_profit],["Net cash flow",summary.cash_flow.net_cash_flow],["Current inventory cost",summary.inventory_valuation.purchase_value]])
        sales=self.limited(self.db.query(Sale).options(joinedload(Sale.items)).filter(Sale.store_id==store_id,Sale.sale_date>=start,Sale.sale_date<end,Sale.status.notin_([SaleStatus.CANCELLED,SaleStatus.VOIDED])))
        sheet("Sales invoices",["Invoice","Date","Customer","Payment","Reference","Subtotal","Bill discount","Total","Status"],[[s.invoice_number,s.sale_date.isoformat(),s.customer_name,s.payment_mode,s.payment_reference,s.subtotal,s.discount_amount,s.total_amount,s.status.value] for s in sales])
        sheet("Recorded GST sales",["Invoice","Date","Product","SKU","Size","HSN","GST rate","Qty","Taxable","CGST","SGST","IGST"],[[s.invoice_number,s.sale_date.isoformat(),i.product_name,i.sku_snapshot,i.size_snapshot,i.hsn_snapshot,i.gst_rate_snapshot,i.quantity,i.taxable_value,i.cgst_amount,i.sgst_amount,i.igst_amount] for s in sales for i in s.items])
        returns=self.limited(self.db.query(SaleReturn,SaleReturnItem,SaleItem,Sale).select_from(SaleReturnItem).join(SaleReturn,SaleReturn.id==SaleReturnItem.sale_return_id).join(SaleItem,SaleItem.id==SaleReturnItem.sale_item_id).join(Sale,Sale.id==SaleReturn.sale_id).filter(SaleReturn.store_id==store_id,SaleReturn.created_at>=start,SaleReturn.created_at<end))
        sheet("Customer returns",["Return","Original invoice","Date","Product","Size","Qty","Sellable","Refund method","Refund","HSN","Original GST rate"],[[str(r.id),s.invoice_number,r.created_at.isoformat(),i.product_name,i.size_snapshot,line.quantity,line.restock,r.refund_method,line.refund_amount,i.hsn_snapshot,i.gst_rate_snapshot] for r,line,i,s in returns])
        purchases=self.limited(self.db.query(Purchase).filter(Purchase.store_id==store_id,Purchase.purchase_date>=start_date,Purchase.purchase_date<=end_date,Purchase.status==PurchaseStatus.CONFIRMED))
        sheet("Purchases",["Invoice","Date","Supplier","Total","Amount paid","Payment"],[[p.invoice_number,str(p.purchase_date),p.supplier_name,p.total_amount,p.amount_paid,p.payment_mode] for p in purchases])
        credits=self.limited(self.db.query(PurchaseReturn).filter(PurchaseReturn.store_id==store_id,PurchaseReturn.created_at>=start,PurchaseReturn.created_at<end))
        sheet("Supplier credits",["Return","Purchase","Date","Credit note","Reason","Credit"],[[str(r.id),str(r.purchase_id),r.created_at.isoformat(),r.credit_note,r.reason,r.credit_amount] for r in credits])
        movements=self.limited(self.db.query(StockHistory).options(joinedload(StockHistory.product)).filter(StockHistory.store_id==store_id,StockHistory.movement_date>=start,StockHistory.movement_date<end))
        sheet("Stock movements",["Date","Product","Variant","Type","Qty","Before","After","Reference"],[[m.movement_date.isoformat(),m.product.name,str(m.product_variant_id),m.movement_type.value,m.qty,m.before_stock,m.after_stock,m.reference] for m in movements])
        variants=self.limited(self.db.query(ProductVariant).options(joinedload(ProductVariant.product).joinedload(Product.category),joinedload(ProductVariant.product).joinedload(Product.brand)).filter(ProductVariant.store_id==store_id))
        stock_headers=["Product","Category","Brand","Size","Colour","SKU","Barcode","Stock","Minimum"]
        sheet("Current inventory",stock_headers, [list(self.stock_row(v).values()) for v in variants])
        sheet("Current low stock",stock_headers,[list(self.stock_row(v).values()) for v in self.limited(self.low_stock_query(user))])
        expenses=self.limited(self.db.query(Expense).filter(Expense.store_id==store_id,Expense.expense_date>=start_date,Expense.expense_date<=end_date))
        sheet("Expenses",["Date","Description","Payment","Amount"],[[str(e.expense_date),e.title,e.payment_mode,e.amount] for e in expenses])
        closings=self.limited(self.db.query(DayClosing).filter(DayClosing.store_id==store_id,DayClosing.business_date>=start_date,DayClosing.business_date<=end_date))
        sheet("Day closing",["Date","Opening","Expected","Counted","Difference","Status","Notes"],[[str(c.business_date),c.opening_cash,c.expected_cash,c.actual_cash,c.difference,c.status,c.notes] for c in closings])
        payments={}
        for s in sales:payments[s.payment_mode]=payments.get(s.payment_mode,Decimal("0"))+s.total_amount
        for r,line,_,_ in returns:payments[r.refund_method]=payments.get(r.refund_method,Decimal("0"))-line.refund_amount
        sheet("Payment modes",["Method","Sales less refunds"],list(payments.items()))
        customers=self.limited(self.db.query(Customer).filter(Customer.store_id==store_id))
        sheet("Current receivables",["Customer","Phone","Balance due"],[[c.name,c.phone,CustomerService(self.db)._totals(c,store_id)[2]] for c in customers])
        suppliers=self.limited(self.db.query(Supplier).filter(Supplier.store_id==store_id))
        sheet("Current payables",["Supplier","Balance due"],[[s.name,SupplierService(self.db)._totals(s,store_id)[2]] for s in suppliers])
        sheet("Notes",["Scope","Explanation"],[["Sales","Original invoices; customer returns are separate adjustment rows."],["GST","Recorded invoice tax details, not a filed tax return. Historical missing tax snapshots are blank/zero. Review return adjustments separately."],["Inventory and balances","Current snapshot, not an as-of-period historical balance."],["Low stock","Per-size threshold overrides the product threshold when configured."]])
        output=BytesIO();book.save(output);return output.getvalue()
