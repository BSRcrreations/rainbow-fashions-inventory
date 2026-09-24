"""Real PostgreSQL proofs for complete retail transactions and safety boundaries."""
import asyncio
import json
from datetime import datetime, timezone, timedelta
from decimal import Decimal
from io import BytesIO
from uuid import uuid4

import pytest
from fastapi import HTTPException, UploadFile
from fastapi.testclient import TestClient
from sqlalchemy import func

from app.core.config import Settings
from app.models import Store, User, Category, Brand, SubCategory, Product, ProductVariant, ProductInventory, StockHistory, Sale, SaleReturn, Purchase, InventoryCostLot
from app.models.enums import UserRole, PricingType, StockScanMode, StockScanStatus
from app.models.operations import OperationsAudit
from app.models.purchase_return import PurchaseReturn
from app.schemas.sale import SaleCreate, SaleReturnCreate, SaleExchangeCreate, SaleVoidRequest
from app.schemas.purchase import QuickPurchaseCreate, PurchaseReturnCreate
from app.schemas.stock_scan import StockScanSessionCreate, QuickStockItemSet, StockScanConfirmRequest
from app.schemas.operations import DayClosingCreate, StoreSettings
from app.schemas.opening_stock_import import OpeningStockImportConfirm, OpeningStockImportReverse
from app.services.stock_scan_service import StockScanService
from app.services.sale_service import SaleService
from app.services.purchase_service import PurchaseService
from app.services.business_service import SupplierService, CustomerService
from app.services.opening_stock_import_service import OpeningStockImportService
from app.services.inventory_reconciliation_service import InventoryReconciliationService
from app.services.operations_service import OperationsService
from app.services.backup_status_service import BackupStatusService


@pytest.fixture
def shop(pg_db):
    db = pg_db
    store = Store(name="Disposable Retail Test", code=uuid4().hex[:20]); db.add(store); db.flush()
    owner = User(store_id=store.id, full_name="Test Owner", email=f"{uuid4().hex}@example.test", password_hash="unused-test-only", role=UserRole.OWNER)
    db.add(owner)
    category = Category(store_id=store.id, name="Leggings"); db.add(category); db.flush()
    brand = Brand(store_id=store.id, category_id=category.id, name="Test Brand")
    sub = SubCategory(store_id=store.id, category_id=category.id, name="Ankle"); db.add_all([brand,sub]); db.flush()
    product = Product(store_id=store.id, category_id=category.id, subcategory_id=sub.id, brand_id=brand.id, name="Cotton Legging", sku=uuid4().hex, purchase_price=10, selling_price=20, pricing_type=PricingType.OWN_PRICE, current_stock=0, minimum_stock=2)
    db.add(product); db.flush()
    variants=[]
    for size in ["M","L"]:
        variant=ProductVariant(store_id=store.id, product_id=product.id, size=size, color="Black", internal_sku=uuid4().hex, barcode=uuid4().hex, identity_key=uuid4().hex, selling_price=20, mrp=25, current_stock=0, last_purchase_cost=10, average_cost=10)
        db.add(variant); variants.append(variant)
    db.flush(); db.commit()
    return db, owner, product, variants


def stock(shop, count=5):
    db, owner, product, variants=shop
    service=StockScanService(db)
    payload=StockScanSessionCreate(session_id=uuid4(), mode=StockScanMode.DAILY_STOCK)
    session=service.create_session(payload,owner)
    for variant in variants:
        service.set_quick_item(session.id,QuickStockItemSet(product_variant_id=variant.id,quantity=count,unit_cost=10),owner)
    service.confirm(session.id,StockScanConfirmRequest(),owner)
    return session


def sale(shop, quantity=1, **kwargs):
    db, owner, _, variants=shop
    return SaleService(db).create(SaleCreate(payment_mode="CASH",items=[{"product_variant_id":variants[0].id,"quantity":quantity}], **kwargs),owner,idempotency_key=str(uuid4()))


def test_quick_stock_draft_recovery_absolute_retries_exact_size_and_reconciliation(shop):
    db,owner,product,variants=shop; service=StockScanService(db)
    request=StockScanSessionCreate(session_id=uuid4(),mode=StockScanMode.DAILY_STOCK)
    session=service.create_session(request,owner)
    assert service.create_session(request,owner).id==session.id
    line=QuickStockItemSet(product_variant_id=variants[1].id,quantity=3,unit_cost=10)
    service.set_quick_item(session.id,line,owner); service.set_quick_item(session.id,line,owner)
    assert variants[1].current_stock==0
    assert len(service.get_session(session.id,owner).items)==1
    service.confirm(session.id,StockScanConfirmRequest(),owner); service.confirm(session.id,StockScanConfirmRequest(),owner)
    db.expire_all()
    assert [v.current_stock for v in variants]==[0,3]
    assert product.current_stock==3
    assert db.query(StockHistory).filter_by(product_id=product.id).count()==1
    assert all(row.category=="HEALTHY" for row in InventoryReconciliationService(db).report(owner))


def test_sale_retry_negative_stock_and_exact_variant(shop):
    stock(shop,2); db,owner,product,variants=shop; service=SaleService(db)
    payload=SaleCreate(payment_mode="CASH",items=[{"product_variant_id":variants[0].id,"quantity":2}]); key=str(uuid4())
    first=service.create(payload,owner,idempotency_key=key)
    assert service.create(payload,owner,idempotency_key=key).id==first.id
    with pytest.raises(HTTPException): service.create(payload,owner,idempotency_key=str(uuid4()))
    db.expire_all(); assert [v.current_stock for v in variants]==[0,2]; assert product.current_stock==2
    assert db.query(Sale).filter_by(store_id=owner.store_id).count()==1


def test_discounted_damaged_return_does_not_restock_and_retry_is_safe(shop):
    stock(shop); db,owner,product,variants=shop
    bill=sale(shop,2,discount_type="PERCENTAGE",discount_value=10)
    request=SaleReturnCreate(reason="Damaged stitching",refund_method="CASH",items=[{"sale_item_id":bill.items[0].id,"quantity":1,"restock":False}])
    key=str(uuid4()); service=SaleService(db); returned=service.create_return(bill.id,request,owner,key)
    assert returned.refund_amount==Decimal("18.00")
    assert service.create_return(bill.id,request,owner,key).id==returned.id
    db.expire_all(); assert variants[0].current_stock==3


def test_exchange_is_atomic_and_idempotent(shop):
    stock(shop); db,owner,product,variants=shop; bill=sale(shop); service=SaleService(db)
    payload=SaleExchangeCreate(reason="Change size",return_items=[{"sale_item_id":bill.items[0].id,"quantity":1}],items=[{"product_variant_id":variants[1].id,"quantity":1}])
    key=str(uuid4()); result=service.exchange(bill.id,payload,owner,key)
    assert result["amount_due"]==0
    assert service.exchange(bill.id,payload,owner,key)["sale"].id==result["sale"].id
    db.expire_all(); assert [v.current_stock for v in variants]==[5,4]
    assert all(row.category=="HEALTHY" for row in InventoryReconciliationService(db).report(owner))


def test_failed_exchange_rolls_back_return(shop):
    stock(shop,1); db,owner,_,variants=shop; bill=sale(shop)
    payload=SaleExchangeCreate(reason="Change size",return_items=[{"sale_item_id":bill.items[0].id,"quantity":1}],items=[{"product_variant_id":variants[1].id,"quantity":99}])
    with pytest.raises(HTTPException): SaleService(db).exchange(bill.id,payload,owner,str(uuid4()))
    db.expire_all(); assert variants[0].current_stock==0
    assert db.query(SaleReturn).filter_by(sale_id=bill.id).count()==0


def test_quick_purchase_draft_confirm_retry_and_supplier_return(shop):
    from app.schemas.supplier import SupplierCreate
    db,owner,product,variants=shop
    supplier=SupplierService(db).create(SupplierCreate(name="Wholesale"),owner)
    service=PurchaseService(db)
    payload=QuickPurchaseCreate(supplier_id=supplier.id,items=[{"product_variant_id":variants[0].id,"quantity":4,"purchase_cost":10}])
    draft=service.quick_purchase(payload,owner,str(uuid4()))
    assert variants[0].current_stock==0
    purchase=service.confirm(draft.id,owner); service.confirm(draft.id,owner)
    db.expire_all(); assert variants[0].current_stock==4
    returned=service.create_return(purchase.id,PurchaseReturnCreate(reason="Damaged delivery",items=[{"purchase_item_id":purchase.items[0].id,"quantity":1}]),owner,str(uuid4()))
    assert returned.credit_amount==10
    assert SupplierService(db).detail(supplier.id,owner).balance_due==30
    assert any(row.entry_type=="PURCHASE_RETURN" for row in SupplierService(db).detail(supplier.id,owner).ledger)
    db.expire_all(); assert product.current_stock==variants[0].current_stock==3


def test_sale_void_restores_once(shop):
    stock(shop); db,owner,_,variants=shop; bill=sale(shop)
    SaleService(db).void(bill.id,SaleVoidRequest(reason="Wrong bill",version=bill.version),owner)
    db.expire_all(); assert variants[0].current_stock==5
    with pytest.raises(HTTPException): SaleService(db).void(bill.id,SaleVoidRequest(reason="Wrong bill",version=bill.version),owner)
    db.expire_all(); assert variants[0].current_stock==5


def test_day_closing_persist_approve_and_audit_are_store_scoped(shop):
    stock(shop); db,owner,_,_=shop; sale(shop)
    service=OperationsService(db); day=datetime.now(timezone.utc).astimezone().date()
    payload=DayClosingCreate(business_date=day,opening_cash=100,actual_cash=120)
    closing=service.submit_closing(payload,owner,"closing-test")
    assert closing.expected_cash==120 and closing.difference==0
    assert service.approve_closing(closing.id,owner,"approve-test").status=="APPROVED"
    actions={entry.action for entry in service.audit_log(owner).items}
    assert {"DAY_CLOSING_APPROVED","DAY_CLOSING_SUBMITTED"} <= actions


def test_api_cashier_owner_boundaries_and_other_store_lookup(shop):
    from app.main import app
    from app.api.deps import get_current_user
    from app.database.session import get_db
    db,owner,product,variants=shop
    cashier=User(id=uuid4(),store_id=owner.store_id,role=UserRole.CASHIER,full_name="Cashier")
    app.dependency_overrides[get_db]=lambda:db
    app.dependency_overrides[get_current_user]=lambda:cashier
    try:
        client=TestClient(app)
        assert client.get('/api/v1/users').status_code==403
        assert client.post('/api/v1/stock-scan/sessions',json={"mode":"DAILY_STOCK"}).status_code==403
        assert client.get('/api/v1/settings/store').status_code==200
        outsider=User(id=uuid4(),store_id=uuid4(),role=UserRole.OWNER,full_name="Other shop")
        app.dependency_overrides[get_current_user]=lambda:outsider
        assert client.get(f'/api/v1/products/{product.id}').status_code==404
        assert client.get(f'/api/v1/sales/catalog/variant/{variants[0].id}').status_code==404
    finally: app.dependency_overrides.clear()


def test_opening_stock_atomic_validation_idempotency_shared_sizes_and_reversal(shop,tmp_path):
    db,owner,_,_=shop
    service=OpeningStockImportService(db,Settings(app_env="testing",allow_test_opening_stock_import_bypass=True,opening_stock_import_dir=tmp_path))
    csv=b"product_name,category,subcategory,brand,sku,barcode,quantity,purchase_cost,selling_price,size,color\nNew Legging,Women,Ankle,Brand,NEW-M,TEST-SHARED,2,10,20,M,Black\nNew Legging,Women,Ankle,Brand,NEW-L,TEST-SHARED,3,10,20,L,Black\n"
    batch=asyncio.run(service.upload_and_validate(UploadFile(filename="stock.csv",file=BytesIO(csv)),owner,"import-test",True))
    assert batch.error_count==0
    request=OpeningStockImportConfirm(confirmation="POST OPENING STOCK",idempotency_key=str(uuid4()))
    result=service.confirm(batch.id,request,owner,"confirm-test")
    assert result.total_quantity==5 and result.created_variants==2
    assert service.confirm(batch.id,request,owner,"retry").already_completed
    targets=StockScanService(db).shared_barcode_targets("TEST-SHARED",owner)
    assert len(targets)==2 and {t.size for t in targets}=={"M","L"}
    with pytest.raises(HTTPException): StockScanService(db).resolve_barcode("TEST-SHARED",owner)
    service.reverse(batch.id,OpeningStockImportReverse(confirmation="REVERSE OPENING STOCK",reason="Test reversal only"),owner,"reverse-test")
    db.expire_all(); assert sum(t.current_stock for t in db.query(ProductVariant).filter(ProductVariant.internal_sku.in_(["NEW-M","NEW-L"])))==0
    imported_lots = db.query(InventoryCostLot).filter(InventoryCostLot.product_variant_id.in_([target.variant_id for target in targets])).all()
    assert imported_lots and all(lot.remaining_quantity == 0 for lot in imported_lots)
    imported_products = {target.product_id for target in targets}
    findings = [row for row in InventoryReconciliationService(db).report(owner) if row.product_id in imported_products]
    assert findings and all(row.category == "HEALTHY" for row in findings)


def test_opening_stock_invalid_rows_are_retained_without_posting(shop, tmp_path):
    db, owner, _, _ = shop
    service = OpeningStockImportService(db, Settings(app_env="staging", allow_test_opening_stock_import_bypass=True, opening_stock_import_dir=tmp_path))
    csv = b"product_name,category,subcategory,brand,sku,barcode,quantity,purchase_cost,selling_price\nInvalid Legging,Leggings,Ankle,Test Brand,INVALID-S,INVALID-BARCODE,-2,10,20\n"
    batch = asyncio.run(service.upload_and_validate(UploadFile(filename="invalid.csv", file=BytesIO(csv)), owner, "invalid-import"))
    assert batch.status == "REVIEW_REQUIRED"
    assert batch.row_count == 1 and batch.valid_row_count == 0 and batch.error_count == 1
    assert batch.total_quantity == 0
    assert db.query(StockHistory).filter_by(store_id=owner.store_id).count() == 0


def test_backup_proof_rejects_stale_success_and_requires_restore(tmp_path):
    files=["latest-database-backup.json","latest-upload-manifest.json","latest-offsite-backup.json","latest-database-restore-test.json","latest-upload-restore-test.json"]
    for name in files:
        (tmp_path/name).write_text(json.dumps({"status":"SUCCESS","timestamp":datetime.now(timezone.utc).isoformat()}))
    assert BackupStatusService(tmp_path).status().posting_allowed
    (tmp_path/files[0]).write_text(json.dumps({"status":"SUCCESS","timestamp":(datetime.now(timezone.utc)-timedelta(days=3)).isoformat()}))
    assert not BackupStatusService(tmp_path).status().posting_allowed


@pytest.mark.parametrize("money_value",["NaN","Infinity","1.001","-1"])
def test_opening_money_rejects_nonfinite_or_fractional_paise(money_value):
    service=OpeningStockImportService(None)
    row={"product_name":"A","category":"B","subcategory":"C","brand":"D","sku":"E","barcode":"F","quantity":"1","purchase_cost":money_value,"selling_price":"20"}
    assert any(error[1]=="INVALID_MONEY" for error in service._normalize_row(row)[1])


def test_tax_snapshot_and_reports_reflect_returns_without_changing_paid_bill(shop):
    from app.services.report_service import ReportService
    from datetime import date
    stock(shop, 3)
    db, owner, product, variants = shop
    product.gst_rate = Decimal("5"); product.hsn_code = "6115"
    variants[0].selling_price = Decimal("21"); db.commit()
    bill = sale(shop, 2)
    assert bill.total_amount == Decimal("42")
    assert bill.items[0].taxable_value == Decimal("40")
    assert bill.items[0].cgst_amount == Decimal("1")
    product.gst_rate = Decimal("18"); db.commit()
    returned = SaleService(db).create_return(bill.id, SaleReturnCreate(reason="Size returned", items=[{"sale_item_id":bill.items[0].id,"quantity":1}]), owner, str(uuid4()))
    report = ReportService(db).summary(owner, date(2026,1,1), date(2027,1,1))
    assert returned.refund_amount == Decimal("21")
    assert report.profit_and_loss.sales_total == Decimal("21")
    assert report.profit_and_loss.gross_profit == Decimal("11")
    assert report.cash_flow.cash_sales == Decimal("21")
    assert bill.items[0].gst_rate_snapshot == Decimal("5")


def test_dashboard_does_not_include_another_store(shop):
    from app.services.dashboard_service import DashboardService
    stock(shop, 2); db, owner, product, variants = shop
    other = Store(name="Different Store", code=uuid4().hex[:20]); db.add(other); db.flush()
    stranger = User(store_id=other.id, full_name="Other", email=f"{uuid4().hex}@example.test", password_hash="unused", role=UserRole.OWNER); db.add(stranger); db.commit()
    assert DashboardService(db).summary(stranger).total_stock == 0
    assert DashboardService(db).summary(owner).total_stock == 4


def test_multi_size_stock_confirmation_keeps_all_aggregate_totals(shop):
    db, owner, product, variants = shop
    stock(shop, 5)
    db.expire_all()
    assert product.current_stock == 10
    assert db.query(ProductInventory).filter_by(product_id=product.id, store_id=owner.store_id).one().current_stock == 10
    assert all(row.category == "HEALTHY" for row in InventoryReconciliationService(db).report(owner))


def test_twenty_thousand_row_opening_preview_is_paged_and_does_not_post_stock(shop, tmp_path):
    import time
    import csv
    from io import StringIO
    from app.services.opening_stock_import_service import REQUIRED_HEADERS, OPTIONAL_HEADERS
    db, owner, product, variants = shop
    service = OpeningStockImportService(db)
    service.settings = Settings(app_env="testing", allow_test_opening_stock_import_bypass=True, opening_stock_import_dir=tmp_path)
    data = StringIO(); writer = csv.writer(data); writer.writerow(REQUIRED_HEADERS + OPTIONAL_HEADERS)
    for number in range(20000):
        writer.writerow([f"Preview Product {number // 20}", "Women", "Leggings", "Preview Brand", f"PREVIEW-{number}", f"BAR-{number}", "1", "10.00", "20.00", str(number % 20), "Black", "", "25.00", "", "", "", "Each", "Main"])
    started = time.monotonic()
    result = asyncio.run(service.upload_and_validate(UploadFile(filename="capacity-preview.csv", file=BytesIO(data.getvalue().encode())), owner, "capacity-proof"))
    assert result.row_count == result.valid_row_count == 20000
    assert result.total_quantity == 20000
    assert result.error_count == 0
    rows, errors = service.detail_rows(result.id, page=100)
    assert len(rows) == 200 and rows[-1].row_number == 20001 and not errors
    assert db.query(Product).filter_by(store_id=owner.store_id).count() == 1
    assert all(variant.current_stock == 0 for variant in variants)
    assert db.query(StockHistory).filter_by(store_id=owner.store_id).count() == 0
    print(f"20,000-row preview elapsed: {time.monotonic() - started:.2f}s; no stock posted")


@pytest.mark.parametrize("method", ["CASH", "UPI", "CARD", "BANK", "OTHER"])
def test_payment_modes_and_required_reference_are_recorded(shop, method):
    stock(shop, 2); db, owner, _, variants = shop
    OperationsService(db).save_settings(StoreSettings(require_payment_reference=True), owner, "payment-settings")
    service = SaleService(db)
    if method != "CASH":
        with pytest.raises(HTTPException):
            service.create(SaleCreate(payment_mode=method, items=[{"product_variant_id": variants[0].id, "quantity": 1}]), owner, idempotency_key=str(uuid4()))
        assert variants[0].current_stock == 2
    bill = service.create(SaleCreate(payment_mode=method, payment_reference="SYNTHETIC-REFERENCE", items=[{"product_variant_id": variants[0].id, "quantity": 1}]), owner, idempotency_key=str(uuid4()))
    assert bill.payment_mode == method and bill.payment_reference == "SYNTHETIC-REFERENCE"
    assert variants[0].current_stock == 1


def test_credit_limit_and_return_reduce_customer_balance(shop):
    from app.models.customer import Customer
    stock(shop, 4); db, owner, _, variants = shop
    customer = Customer(store_id=owner.store_id, name="Credit fixture", credit_limit=30); db.add(customer); db.commit()
    service = SaleService(db)
    payload = SaleCreate(customer_id=customer.id, payment_mode="CREDIT", items=[{"product_variant_id": variants[0].id, "quantity": 1}])
    bill = service.create(payload, owner, idempotency_key=str(uuid4()))
    with pytest.raises(HTTPException): service.create(payload, owner, idempotency_key=str(uuid4()))
    assert CustomerService(db)._totals(customer, owner.store_id)[2] == 20
    service.create_return(bill.id, SaleReturnCreate(reason="Size returned", refund_method="CREDIT", items=[{"sale_item_id": bill.items[0].id, "quantity": 1}]), owner, str(uuid4()))
    assert CustomerService(db)._totals(customer, owner.store_id)[2] == 0


def test_cashier_combined_discount_limit_cannot_be_bypassed_with_item_discount(shop):
    stock(shop, 2); db, owner, _, variants = shop
    cashier = User(store_id=owner.store_id, full_name="Cashier", email=f"{uuid4().hex}@example.com", password_hash="unused", role=UserRole.CASHIER); db.add(cashier); db.commit()
    payload = SaleCreate(payment_mode="CASH", discount_type="PERCENTAGE", discount_value=10, items=[{"product_variant_id": variants[0].id, "quantity": 1, "discount_type": "PERCENTAGE", "discount_value": 10}])
    with pytest.raises(HTTPException) as error: SaleService(db).create(payload, cashier, idempotency_key=str(uuid4()))
    assert error.value.status_code == 403 and variants[0].current_stock == 2
    assert SaleService(db).create(payload, owner, idempotency_key=str(uuid4())).total_amount == Decimal("16.20")


def test_variant_threshold_and_excel_report_preserve_stock_and_scope(shop):
    from datetime import date
    from openpyxl import load_workbook
    from app.schemas.product import ProductVariantUpdate
    from app.services.variant_management_service import VariantManagementService
    from app.services.report_export_service import ReportExportService
    stock(shop, 2); db, owner, product, variants = shop
    product.minimum_stock=0;product.name="=UNTRUSTED NAME";product.gst_rate=5;db.commit()
    history_count=db.query(StockHistory).count()
    VariantManagementService(db).update(variants[1].id,ProductVariantUpdate(minimum_stock=3),owner,"threshold-test")
    assert db.query(StockHistory).count()==history_count
    assert [v.current_stock for v in variants]==[2,2]
    exporter=ReportExportService(db)
    low=exporter.low_stock(owner)
    assert len(low['items'])==1 and low['items'][0]['size']=='L' and low['items'][0]['minimum_stock']==3
    assert exporter.low_stock(owner,category_id=uuid4())['items']==[]
    bill=sale(shop)
    data=exporter.workbook(owner,date(2026,1,1),date(2026,12,31))
    book=load_workbook(BytesIO(data),data_only=False)
    assert {'Sales invoices','Recorded GST sales','Current low stock','Current receivables','Current payables','Day closing'} <= set(book.sheetnames)
    assert book['Sales invoices'].cell(2,1).value==bill.invoice_number
    assert book['Recorded GST sales'].cell(2,3).value=="'=UNTRUSTED NAME"
    assert book['Recorded GST sales'].cell(2,7).value==5
    other=Store(name='Other report store',code=uuid4().hex[:20]);db.add(other);db.flush()
    outsider=User(id=uuid4(),store_id=other.id,role=UserRole.OWNER,full_name='Other')
    assert exporter.low_stock(outsider)['items']==[]
    other_book=load_workbook(BytesIO(exporter.workbook(outsider,date(2026,1,1),date(2026,12,31))))
    assert other_book['Sales invoices'].max_row==1


def test_certification_exact_four_size_lifecycle_and_shared_barcode(shop, tmp_path):
    from app.schemas.stock_scan import VariantStockStageRequest
    db, owner, _, _ = shop
    service = OpeningStockImportService(db, Settings(app_env="testing", allow_test_opening_stock_import_bypass=True, opening_stock_import_dir=tmp_path))
    csv = 'product_name,category,subcategory,brand,sku,barcode,quantity,purchase_cost,selling_price,size,color\n'
    csv += ''.join(f'Twin Birds Legging,Women,Ankle,Twin Birds,CERT-{size},123456,{quantity},10,20,{size},Black\n' for size, quantity in [('S',10),('M',15),('L',20),('XL',12)])
    batch = asyncio.run(service.upload_and_validate(UploadFile(filename='pilot.csv', file=BytesIO(csv.encode())), owner, 'certification-pilot', True))
    assert batch.error_count == 0
    service.confirm(batch.id, OpeningStockImportConfirm(confirmation='POST OPENING STOCK',idempotency_key=str(uuid4())),owner,'certification-pilot')
    scanner = StockScanService(db)
    targets = scanner.shared_barcode_targets('123456', owner)
    assert {t.size for t in targets} == {'S','M','L','XL'}
    variants = {t.size: db.get(ProductVariant,t.variant_id) for t in targets}
    product = variants['S'].product
    def assert_stock(expected):
        db.expire_all()
        assert {size:variant.current_stock for size,variant in variants.items()} == expected
        total = sum(expected.values())
        assert product.current_stock == total
        assert db.query(ProductInventory).filter_by(product_id=product.id,store_id=owner.store_id).one().current_stock == total
        findings = [row for row in InventoryReconciliationService(db).report(owner) if row.product_id == product.id]
        assert findings and all(row.category == 'HEALTHY' for row in findings), [row.model_dump(mode='json') for row in findings]
    assert_stock({'S':10,'M':15,'L':20,'XL':12})
    with pytest.raises(HTTPException) as ambiguity: SaleService(db).variant_by_barcode('123456',owner)
    assert ambiguity.value.status_code == 409
    bill = SaleService(db).create(SaleCreate(payment_mode='CASH',items=[{'product_variant_id':variants['M'].id,'quantity':2}]),owner,idempotency_key=str(uuid4()))
    assert_stock({'S':10,'M':13,'L':20,'XL':12})
    SaleService(db).create_return(bill.id,SaleReturnCreate(reason='Size return',refund_method='CASH',items=[{'sale_item_id':bill.items[0].id,'quantity':1}]),owner,str(uuid4()))
    assert_stock({'S':10,'M':14,'L':20,'XL':12})
    correction = scanner.create_session(StockScanSessionCreate(mode='STOCK_ADJUSTMENT',notes='Physical recount'),owner)
    scanner.stage_selected_variant(correction.id,VariantStockStageRequest(barcode='123456',product_variant_id=variants['XL'].id,quantity=10,confirm_shared_barcode=True),owner)
    scanner.confirm(correction.id,StockScanConfirmRequest(notes='Physical count difference'),owner)
    assert_stock({'S':10,'M':14,'L':20,'XL':10})
    purchase = PurchaseService(db).quick_purchase(QuickPurchaseCreate(items=[{'product_variant_id':variants['L'].id,'quantity':5,'purchase_cost':10}]),owner,str(uuid4()))
    PurchaseService(db).confirm(purchase.id,owner)
    assert_stock({'S':10,'M':14,'L':25,'XL':10})
    with pytest.raises(HTTPException): service.reverse(batch.id,OpeningStockImportReverse(confirmation='REVERSE OPENING STOCK',reason='Conflicting later activity'),owner,'blocked-reversal')


def test_credit_exchange_reuses_credit_allowance_and_rolls_back_failed_replacement(shop):
    from app.models.customer import Customer
    stock(shop,4); db,owner,_,variants = shop
    customer=Customer(store_id=owner.store_id,name='Credit exchange',credit_limit=20); db.add(customer); db.commit()
    service=SaleService(db)
    bill=service.create(SaleCreate(customer_id=customer.id,payment_mode='CREDIT',items=[{'product_variant_id':variants[0].id,'quantity':1}]),owner,idempotency_key=str(uuid4()))
    payload=SaleExchangeCreate(payment_mode='CREDIT',reason='Change size',return_items=[{'sale_item_id':bill.items[0].id,'quantity':1}],items=[{'product_variant_id':variants[1].id,'quantity':2}])
    with pytest.raises(HTTPException): service.exchange(bill.id,payload,owner,str(uuid4()))
    assert db.query(SaleReturn).filter_by(sale_id=bill.id).count() == 0
    payload.items[0].quantity=1; key=str(uuid4())
    result=service.exchange(bill.id,payload,owner,key)
    assert service.exchange(bill.id,payload,owner,key)['sale'].id == result['sale'].id
    assert result['sale'].exchange_return_id == result['sale_return'].id
    assert CustomerService(db)._totals(customer,owner.store_id)[2] == 20
    db.expire_all(); assert [v.current_stock for v in variants] == [4,3]


def test_formal_purchase_atomic_save_retry_stale_version_and_failed_line(shop):
    from app.schemas.purchase import PurchaseDraftSave
    db,owner,product,variants=shop; service=PurchaseService(db)
    draft=service.quick_purchase(QuickPurchaseCreate(items=[{'product_variant_id':variants[0].id,'quantity':1,'purchase_cost':10}]),owner,str(uuid4()))
    payload=PurchaseDraftSave(header={'version':draft.version,'invoice_number':'CERT-FORMAL','notes':'saved together'},items=[{'product_id':product.id,'product_variant_id':variants[1].id,'product_name':product.name,'size':'L','color':'Black','quantity':3,'purchase_price':10,'line_total':30,'user_verified':True}])
    key=str(uuid4()); saved=service.save_complete_draft(draft.id,payload,owner,key)
    assert len(saved.items)==1 and saved.items[0].quantity==3
    assert service.save_complete_draft(draft.id,payload,owner,key).version==saved.version
    with pytest.raises(HTTPException): service.save_complete_draft(draft.id,payload,owner,str(uuid4()))
    payload.header.version=saved.version; payload.items[0].product_variant_id=uuid4(); payload.header.notes='must roll back'
    with pytest.raises(HTTPException): service.save_complete_draft(draft.id,payload,owner,str(uuid4()))
    db.expire_all(); current=service.get(draft.id,owner)
    assert current.notes=='saved together' and len(current.items)==1
    assert product.current_stock==0 and db.query(StockHistory).filter_by(product_id=product.id).count()==0
    service.confirm(draft.id,owner); db.expire_all()
    assert [v.current_stock for v in variants]==[0,3]


def test_formal_purchase_exact_barcode_reuses_existing_product_and_variant(shop):
    from app.schemas.purchase import PurchaseDraftSave
    db,owner,product,variants=shop; service=PurchaseService(db)
    draft=service.quick_purchase(QuickPurchaseCreate(items=[{'product_variant_id':variants[0].id,'quantity':1,'purchase_cost':10}]),owner,str(uuid4()))
    payload=PurchaseDraftSave(
        header={'version':draft.version,'invoice_number':'CERT-FORMAL-EXACT-ID'},
        items=[{
            'product_name':product.name,
            'category_id':product.category_id,
            'brand_id':product.brand_id,
            'barcode':variants[1].barcode,
            'internal_sku':variants[1].internal_sku,
            'size':'L',
            'color':'Black',
            'quantity':2,
            'purchase_price':10,
            'line_total':20,
            'user_verified':True,
        }],
    )
    saved=service.save_complete_draft(draft.id,payload,owner,str(uuid4()))
    service.confirm(saved.id,owner)
    db.expire_all()
    assert db.query(Product).filter(Product.store_id==owner.store_id,Product.name==product.name).count()==1
    assert db.query(ProductVariant).filter(ProductVariant.product_id==product.id).count()==2
    assert [variant.current_stock for variant in variants]==[0,2]


@pytest.mark.parametrize('role', ['MANAGER','CASHIER','STOCK_STAFF','ACCOUNTANT','VIEWER'])
def test_all_nonowners_denied_direct_owner_apis(shop,role):
    from app.main import app
    from app.api.deps import get_current_user
    from app.database.session import get_db
    db,owner,_,_=shop
    actor=User(id=uuid4(),store_id=owner.store_id,role=UserRole(role),full_name='Role test')
    app.dependency_overrides[get_db]=lambda:db
    app.dependency_overrides[get_current_user]=lambda:actor
    try:
        client=TestClient(app)
        for method,path,body in [('GET','/users',None),('PUT','/settings/store',{}),('POST',f'/opening-stock-imports/{uuid4()}/reverse',{'confirmation':'REVERSE OPENING STOCK','reason':'forbidden test'})]:
            response=client.request(method,'/api/v1'+path,json=body)
            assert response.status_code==403,(role,path,response.status_code)
        if role=='VIEWER':
            assert client.post('/api/v1/sales',json={'payment_mode':'CASH','items':[]}).status_code==403
    finally: app.dependency_overrides.clear()


def test_shop_dashboard_accounts_thresholds_and_net_payment_collection(shop):
    from app.services.shop_summary_service import shop_summary
    from app.models.customer import Customer
    db,owner,product,variants=shop; stock(shop,3)
    variants[0].minimum_stock=4; db.commit()
    customer=Customer(store_id=owner.store_id,name='Account',opening_credit=7); db.add(customer); db.commit()
    bill=sale(shop)
    SaleService(db).create_return(bill.id,SaleReturnCreate(reason='Returned item',refund_method='CASH',items=[{'sale_item_id':bill.items[0].id,'quantity':1}]),owner,str(uuid4()))
    summary=shop_summary(db,owner)
    assert summary['customer_receivable']==7
    assert summary['low_stock_count']==1 and summary['out_of_stock_count']==0
    assert summary['inventory_retail_value']==120
    from app.services.sale_service import BUSINESS_TIMEZONE
    day=datetime.now(BUSINESS_TIMEZONE).date()
    collection=SaleService(db)._collection(day,day,owner.store_id)
    assert collection['cash']==0 and collection['total']==0
    foreign=User(store_id=uuid4())
    other=shop_summary(db,foreign)
    assert other['customer_receivable']==0 and other['inventory_retail_value']==0


def test_stock_adjustment_requires_exact_size_and_retries_once_including_zero_count(shop):
    from app.schemas.stock import StockAdjustmentCreate
    from app.services.stock_service import StockService
    stock(shop,3);db,owner,product,variants=shop;service=StockService(db)
    with pytest.raises(HTTPException): service.adjust(StockAdjustmentCreate(product_id=product.id,qty=1,direction='INCREASE',reference='No size selected'),owner)
    payload=StockAdjustmentCreate(product_id=product.id,product_variant_id=variants[0].id,qty=0,direction='DECREASE',adjustment_type='SET_COUNTED_QUANTITY',reference='Physical shelf count')
    key=str(uuid4());first=service.adjust(payload,owner,key)
    assert service.adjust(payload,owner,key).id==first.id
    db.expire_all();assert [v.current_stock for v in variants]==[0,3] and product.current_stock==3
    assert all(row.category=='HEALTHY' for row in InventoryReconciliationService(db).report(owner))


def test_invoice_upload_process_and_review_materialize_lines_without_posting(shop,tmp_path,monkeypatch):
    from reportlab.pdfgen import canvas
    from starlette.datastructures import Headers
    import app.services.purchase_document_service as documents
    import app.services.file_service as files
    db,owner,product,variants=shop
    monkeypatch.setattr(files,'get_settings',lambda:Settings(app_env='testing',upload_dir=tmp_path))
    from app.ai.local_ocr import LocalOCRService
    monkeypatch.setattr(documents,'get_ocr_service',LocalOCRService)
    class BorrowedSession:
        def __getattr__(self,name):return getattr(db,name)
        def close(self):pass
    monkeypatch.setattr(documents,'SessionLocal',BorrowedSession)
    path=tmp_path/'actual-invoice.pdf';pdf=canvas.Canvas(str(path))
    for n,line in enumerate(['Supplier: Test Wholesale','Invoice Number: DOC-CERT-1','Date: 2026-09-13','Test Brand | Leggings | Cotton Legging | M | Black | 2 | 10 | 25 | 20']):pdf.drawString(20,750-n*24,line)
    pdf.save()
    upload=UploadFile(filename=path.name,file=BytesIO(path.read_bytes()),headers=Headers({'content-type':'application/pdf'}))
    document,job,created=asyncio.run(documents.PurchaseDocumentService(db).upload(upload,owner))
    assert created
    documents.PurchaseDocumentService.process(job.id)
    db.expire_all(); ready=documents.PurchaseDocumentService(db).get_job(job.id,owner)
    assert ready.status.value=='REVIEW_REQUIRED',(ready.error_code,ready.error_message)
    result=PurchaseService(db).create_from_document(job.id,owner)
    assert len(result.purchase.items)==1 and result.purchase.items[0].quantity==2
    assert result.purchase.items[0].size=='M'
    assert PurchaseService(db).create_from_document(job.id,owner).purchase.id==result.purchase.id
    assert product.current_stock==0 and db.query(StockHistory).filter_by(store_id=owner.store_id).count()==0
